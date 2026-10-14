"""
Builds the Phase-1 cheque-bounce `.docx` templates (docxtpl-tagged Word
files) into app/drafting/tracks/cheque_bounce/templates/.

Why a script instead of committing the .docx files directly: a `.docx` is a
zipped XML blob — unreviewable in a diff and easy to corrupt by hand. This
script is the reviewable source; running it regenerates the actual template
files. A lawyer reviewing wording should edit the regenerated .docx directly
in Word (docxtpl tags are plain text, safe to leave alone while editing
prose around them) — this script only needs to be re-run if the STRUCTURE
changes (new field, new conditional section).

Usage:
    pip install -r requirements.txt   # adds docxtpl, python-docx
    python3 scripts/build_drafting_templates.py

Builds `legal_notice.docx` now (the Phase-1 vertical slice — LLD S.13,
"start with M1-M3 on the legal notice alone"). Other Phase-1 documents
(complaint, affidavit, annexure index, vakalatnama, synopsis, BSA S.63
certificate) are M4 and intentionally not built by this script yet; add a
`build_<name>()` function here per LLD S.4 / track.yaml when they are.

Formatting is a general, lawyer-reviewable default (Times New Roman 12pt,
1.5 line spacing, 1-inch margins) — NOT a specific court's circular. See LLD
S.14.B: per-court format rules are an open input this spec deliberately
does not invent.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

OUT_DIR = Path(__file__).resolve().parent.parent / "app" / "drafting" / "tracks" / "cheque_bounce" / "templates"


def _set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(8)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)


def _p(doc: Document, text: str = "", *, bold=False, italic=False, underline=False,
       align=None, size=None, color=None, space_after=None):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    if space_after is not None:
        para.paragraph_format.space_after = Pt(space_after)
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return para


def build_legal_notice() -> Path:
    doc = Document()
    _set_base_style(doc)

    # Draft banner — removed once the lawyer finalises (docx_renderer.py
    # passes draft_banner=False at that point).
    _p(
        doc,
        "{% if draft_banner %}DRAFT — FOR ADVOCATE REVIEW, NOT FOR SERVICE{% endif %}",
        bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=(192, 0, 0), size=11,
    )

    # Advocate letterhead block, right-aligned.
    _p(doc, "{{ advocate.name }}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "Advocate", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "{{ advocate.address }}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "Mob: {{ advocate.mobile }}   Email: {{ advocate.email }}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(
        doc,
        "Enrolment No.: {{ advocate.enrolment_no }}, {{ advocate.state_bar_council }}",
        align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=16,
    )

    _p(doc, "Date: {{ today_long }}", space_after=16)

    _p(doc, "To,")
    _p(doc, "{{ accused.name }}")
    _p(doc, "{% if is_corporate_accused %}Through its Director(s)/authorised signatory{% endif %}")
    _p(doc, "{{ accused.address }}", space_after=16)

    _p(
        doc,
        "NOTICE UNDER SECTION 138 OF THE NEGOTIABLE INSTRUMENTS ACT, 1881",
        bold=True, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16,
    )

    _p(doc, "Dear Sir/Madam,")

    _p(
        doc,
        "Under instructions from and on behalf of my client, {{ complainant.name }}, "
        "{% if complainant.kind == 'individual' %}son/daughter of {{ complainant.parentage_or_designation }}, "
        "residing at{% else %}having its registered office/place of business at{% endif %} "
        "{{ complainant.address }} (hereinafter referred to as 'my client'), I address you, "
        "{{ accused.name }}, as follows:",
    )

    _p(doc, "1.\t{{ transaction.narrative }}")

    _p(
        doc,
        "2.\tIn part/full discharge of the aforesaid liability, you issued Cheque No. "
        "{{ relied_cheque.number }} dated {{ cheque_date_long }} drawn on {{ relied_cheque.drawer_bank }}, "
        "{{ relied_cheque.drawer_branch }} for a sum of Rs. {{ cheque_amount_words }}, in favour of my client.",
    )

    _p(
        doc,
        "3.\tThe said cheque, upon presentment for collection through my client's banker at "
        "{{ relied_presentment.presenting_branch }}, was returned unpaid/dishonoured vide return "
        "memo No. {{ relied_dishonour.return_memo_no }} dated {{ relied_dishonour.return_memo_date_long }}, "
        "for the reason '{{ relied_dishonour.reason_text }}'.",
    )

    _p(
        doc,
        "4.\t{% if is_corporate_accused %}You, {{ accused.name }}, are a company/firm, and the "
        "undersigned's client is informed that the following person(s) are in charge of, and "
        "responsible to the company/firm for the conduct of, its business, and are accordingly "
        "liable under Section 141 of the Negotiable Instruments Act, 1881: "
        "{% for d in directors_in_charge %}{{ d.name }} ({{ d.designation }}){% if not loop.last %}, {% endif %}"
        "{% endfor %}.{% endif %}",
    )

    _p(
        doc,
        "5.\tThe aforesaid cheque, having been dishonoured, you are liable to pay my client the "
        "amount thereof together with interest and costs, and you are hereby called upon to pay to "
        "my client the sum of Rs. {{ amount_demanded_words }} within 15 (fifteen) days of the receipt "
        "of this notice, failing which my client shall be constrained to initiate criminal proceedings "
        "against you under Section 138 of the Negotiable Instruments Act, 1881, in addition to civil "
        "proceedings for recovery of the said amount together with interest and costs, entirely at "
        "your risk, cost and consequences.",
    )

    _p(
        doc,
        "6.\tThis notice is issued without prejudice to any other right or remedy available to my "
        "client under law, all of which are expressly reserved.",
        space_after=24,
    )

    _p(doc, "Yours faithfully,")
    _p(doc, "{{ advocate.name }}")
    _p(doc, "Advocate for the Complainant", space_after=24)

    _p(
        doc,
        "{% if draft_banner %}DRAFT — Template legal_notice v{{ template_version }} · "
        "facts {{ facts_hash_short }} · for advocate review before dispatch{% endif %}",
        italic=True, size=9, color=(128, 128, 128),
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / "legal_notice.docx"
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    path = build_legal_notice()
    print(f"Wrote {path}")
